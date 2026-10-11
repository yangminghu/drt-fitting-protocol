"""
【已被 `DRT归档工具箱.py` 取代 · v1.7】新批次请用工具箱的 `summary` 子命令：
    python DRT归档工具箱.py summary "<根>/02_拟合结果_DRT格式" "EIS-DRT <STEM>" --out-dir "<根>/04_汇总表"
本文件与工具箱功能等价（已实测：同一输入产出逐单元格一致），保留仅供复现历史批次；请勿在此就地修改。

生成 DRT 结果汇总表（协议 §8.4）：
  首列 tau，其余每列一个数据集的 γ
  默认输出 .xlsx（DRT / Summary / Info 三工作表）

轴模式:
  --mode union        (默认) 取全体有效窗口的并集 —— 不损失任何频段。
                      每个数据集只在自己窗口内有值，窗口外留【空白】。
  --mode intersection 取交集 —— 得到完整矩阵（无空白），但会砍掉频段较宽
                      的数据集多出来的部分。仅在需要完整矩阵做统计时使用。

用法:
  python DRT汇总表_生成脚本.py <结果目录> <输出前缀> [--mode union|intersection]
                              [--files A.txt,B.txt,...] [--out-dir <目录>] [--csv] [--note "..."]
  （规范结构下 <结果目录> = 02_拟合结果_DRT格式/，--out-dir = 04_汇总表/；
    残差表在 03_残差表/ 会被自动找到 —— 找不到时窗口会退化为拟合网格，
    H14 第 2/3 项随即失效，因此残差表是必需输入，不是可选输入。）
"""
import sys, os, glob, csv, re, argparse
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

# H19：Info 页记录的协议版本必须与实际遵循的版本一致。
# 本常量是唯一来源；协议文档升版时与此处同步修改（禁止在别处硬编码版本字符串）。
PROTOCOL = "DRT 拟合原则 v1.6, §8.4（含 §8.0 命名与目录分类 H16/H17/H18；H19/H20）"

# H14 第 2/3 项的窗口边界判据。生成器与校验器**必须逐字相同**，
# 否则并集轴端点会被判成「窗口内漏空」(GAP) 或「窗口外有值」，产生 12 列全挂的假失败。
BOUND_TOL = 1e-12


def natural_key(p):
    nums = re.findall(r"\d+", os.path.basename(p))
    return [int(x) for x in nums] or [0]


def read_fitted(path):
    L = R = None; tau = []; gam = []
    for line in open(path, encoding="utf-8-sig"):
        s = line.strip()
        if not s:
            continue
        if s.startswith("L,"):
            L = float(s.split(",")[1])
        elif s.startswith("R,"):
            R = float(s.split(",")[1])
        elif s.lower().startswith("tau"):
            continue
        else:
            a, b = s.split(","); tau.append(float(a)); gam.append(float(b))
    tau = np.asarray(tau); gam = np.asarray(gam)
    if np.any(np.diff(tau) <= 0):
        o = np.argsort(tau); tau, gam = tau[o], gam[o]
    name = os.path.basename(path).replace("_fitted_DRT.txt", ".txt")
    return dict(name=name, tau=tau, gam=gam, R=R, L=L)


def drt_stats(it):
    """§8.3 所需的 DRT 侧统计量（全部由 tau/gamma 直接算出）"""
    g, t = it["gam"], it["tau"]
    dln = np.gradient(np.log(t))
    out = {}
    out["R_pol_total"] = float(np.sum(g*dln))
    if it.get("fmin"):
        lo, hi = 1/(2*np.pi*it["fmax"]), 1/(2*np.pi*it["fmin"])
        win = (t >= lo) & (t <= hi)
        out["R_pol_win"] = float(np.sum(g[win]*dln[win]))
    else:
        out["R_pol_win"] = out["R_pol_total"]
    out["R_pol_out"] = out["R_pol_total"] - out["R_pol_win"]
    out["min_gamma"] = float(g.min())
    out["edge_fast"] = float(g[0])
    out["edge_slow"] = float(g[-1])
    return out


def read_residuals(path):
    if not os.path.exists(path):
        return None
    # §8.0.4 H18 要求 csv 带 BOM；必须用 utf-8-sig，否则首列名会变成 '\ufefff_Hz'
    return np.genfromtxt(path, delimiter=",", names=True, encoding="utf-8-sig")


def find_residuals(fitted_dir, base):
    """§8.0.3：残差表可能在同一个目录，也可能在 <根>/03_残差表/"""
    cands = [os.path.join(fitted_dir, base + "_residuals.csv"),
             os.path.join(os.path.dirname(os.path.normpath(fitted_dir)), "03_残差表",
                          base + "_residuals.csv")]
    for c in cands:
        d = read_residuals(c)
        if d is not None:
            return d
    return None


def span_dec(tau):
    return float(np.log10(tau.max()/tau.min()))


def build_axis(items, mode):
    """returns (axis, windows, exact, note)   windows[i] = (tau_lo, tau_hi) of dataset i"""
    los, his = [], []
    for it in items:
        los.append(it["tau"].min() if it["fmin"] is None else 1/(2*np.pi*it["fmax"]))
        his.append(it["tau"].max() if it["fmin"] is None else 1/(2*np.pi*it["fmin"]))
    lo = max(los) if mode == "intersection" else min(los)
    hi = min(his) if mode == "intersection" else max(his)
    if lo >= hi:
        raise SystemExit("ERROR: no common tau window (%s mode) -> check §2 H1" % mode)

    grids = [it["tau"] for it in items]
    identical = all(g.shape == grids[0].shape and np.allclose(g, grids[0], rtol=0, atol=0)
                    for g in grids)
    if identical:
        axis = grids[0][(grids[0] >= lo) & (grids[0] <= hi)]
        note = ("%s axis over [%.6g, %.6g] s; all datasets share one grid -> "
                "its own grid points used, NO interpolation" % (mode, lo, hi))
        return axis, list(zip(los, his)), True, note

    dens = max(len(it["tau"]) / span_dec(it["tau"]) for it in items)
    n = max(2, int(round(np.log10(hi/lo) * dens)))
    axis = np.logspace(np.log10(lo), np.log10(hi), n)
    note = ("%s axis over [%.6g, %.6g] s; %d points at the finest available density "
            "(%.1f pts/decade) -> linear interpolation in ln(tau)"
            % (mode, lo, hi, n, dens))
    return axis, list(zip(los, his)), False, note


def build_matrix(items, axis, windows, tol=BOUND_TOL):
    """NaN wherever a dataset has no value (outside its own measured window).
    边界判据必须与 H14 校验器完全一致，否则端点会被误判为 gap。"""
    M = len(items)
    data = np.full((len(axis), M), np.nan)
    for i, (it, (lo, hi)) in enumerate(zip(items, windows)):
        inside = (axis >= lo*(1 - tol)) & (axis <= hi*(1 + tol))
        if not inside.any():
            continue
        vals = np.interp(np.log(axis[inside]), np.log(it["tau"]), it["gam"],
                         left=np.nan, right=np.nan)
        data[inside, i] = np.maximum(vals, 0.0)
    return data


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("dir"); ap.add_argument("prefix")
    ap.add_argument("--files", default=None)
    ap.add_argument("--note", default="")
    ap.add_argument("--mode", choices=["union", "intersection"], default="union")
    ap.add_argument("--out-dir", default=None,
                    help="输出目录；默认与拟合结果同目录。规范结构下应为 04_汇总表/")
    ap.add_argument("--csv", action="store_true",
                    help="additionally write a .csv twin (off by default)")
    a = ap.parse_args()
    out_dir = a.out_dir or a.dir

    if a.files:
        paths = [os.path.join(a.dir, x.strip()) for x in a.files.split(",")]
    else:
        paths = sorted(glob.glob(os.path.join(a.dir, "*_fitted_DRT.txt")), key=natural_key)
    if not paths:
        raise SystemExit("no fitted files found in " + a.dir)

    items = []
    for p in paths:
        it = read_fitted(p)
        base = os.path.basename(p).replace("_fitted_DRT.txt", "")
        res = find_residuals(a.dir, base)
        if res is not None:
            it["fmin"] = float(res["f_Hz"].min()); it["fmax"] = float(res["f_Hz"].max())
            it["maxRe"] = float(np.abs(res["dRe_percent"]).max())
            it["maxIm"] = float(np.abs(res["dIm_percent"]).max())
            it["rms_pct"] = float(np.sqrt(np.mean(np.concatenate(
                [res["dRe_percent"], res["dIm_percent"]])**2)))
            it["npt"] = int(len(res))
            # H4 判据要用的 min Z'（来自实测列，不是拟合列）
            it["minZre"] = float(np.min(res["Zre_exp"])) if "Zre_exp" in res.dtype.names else None
        else:
            it["fmin"] = it["fmax"] = None
            it["maxRe"] = it["maxIm"] = it["rms_pct"] = float("nan"); it["npt"] = 0
            it["minZre"] = None
        it.update(drt_stats(it))
        items.append(it)

    axis, windows, exact, note = build_axis(items, a.mode)
    data = build_matrix(items, axis, windows)
    M = len(items)
    filled = int(np.isfinite(data).sum())
    print("mode        : %s" % a.mode)
    print("axis        : %d points, tau %.6g .. %.6g s" % (len(axis), axis.min(), axis.max()))
    print("datasets    : %d" % M)
    print("assembly    : %s" % note)
    print("cells       : %d total, %d filled, %d blank (%.1f%% blank)"
          % (data.size, filled, data.size - filled, 100.0*(data.size-filled)/data.size))
    for i, it in enumerate(items):
        lo, hi = windows[i]
        col = data[:, i]
        bad = np.isfinite(col) & ((axis < lo*(1 - BOUND_TOL)) | (axis > hi*(1 + BOUND_TOL)))
        if bad.any():
            print("  WARNING: %s has %d out-of-window values" % (it["name"], bad.sum()))

    if a.csv:
        csv_path = os.path.join(out_dir, a.prefix + ".csv")
        # H18：换行必须是 \n。csv.writer 的默认行终止符是 \r\n，
        # 只写 newline="" 并不够 —— 必须显式 lineterminator="\n"，否则 csv 孪生文件违反 H18。
        with open(csv_path, "w", encoding="utf-8-sig", newline="") as fh:
            w = csv.writer(fh, lineterminator="\n")
            w.writerow(["tau"] + [it["name"] for it in items])
            for r in range(len(axis)):
                w.writerow(["%.12g" % axis[r]] +
                           [("" if not np.isfinite(v) else "%.12g" % v) for v in data[r]])
        print("saved", csv_path)

    xlsx_path = os.path.join(out_dir, a.prefix + ".xlsx")
    wb = Workbook()
    ws = wb.active; ws.title = "DRT"
    ws.append(["tau"] + [it["name"] for it in items])
    for c in range(1, M + 2):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True); cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for r in range(len(axis)):
        ws.cell(row=r + 2, column=1, value=float(axis[r]))
        for c in range(M):
            v = data[r, c]
            if np.isfinite(v):
                ws.cell(row=r + 2, column=c + 2, value=float(v))
            # else: leave the cell EMPTY (never write 0)
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["A"].number_format = "0.0000E+00"
    for c in range(M):
        L = get_column_letter(c + 2)
        ws.column_dimensions[L].width = 13
        for r in range(len(axis)):
            if ws.cell(row=r + 2, column=c + 2).value is not None:
                ws.cell(row=r + 2, column=c + 2).number_format = "0.0000E+00"
    ws.freeze_panes = "B2"

    ws2 = wb.create_sheet("Summary")
    # §8.3 固定列（名字与顺序不得增删），再附 tau 有效窗口
    ws2.append(["sample", "cycle_index", "n_points", "f_min_Hz", "f_max_Hz", "R_ohm", "L_H",
                "R_pol_total_ohm", "R_pol_in_window_ohm", "R_pol_outside_window_ohm",
                "max_dRe_percent", "max_dIm_percent", "rms_percent", "min_gamma",
                "edge_fast_gamma", "edge_slow_gamma", "R_le_minZprime",
                "tau_window_min_s", "tau_window_max_s"])
    for c in range(1, 20):
        ws2.cell(row=1, column=c).font = Font(bold=True)
    for i, it in enumerate(items):
        lo, hi = windows[i]
        # H4 必须真算：R <= min Z'（实测列）。**禁止**用 `R <= 某常数` 之类的占位判据，
        # 那会让这一列恒为 NO，而交付物看上去"有一列布尔结论"，H14 也不会发现。
        if it["R"] is not None and it["minZre"] is not None:
            h4 = "YES" if it["R"] <= it["minZre"] else "NO"
        else:
            h4 = "NA"
        if h4 == "NO":
            print("  WARNING: H4 FAIL for %s: R=%.6g > min Z'=%.6g"
                  % (it["name"], it["R"], it["minZre"]))
        ws2.append([it["name"], i + 1, it["npt"],
                    float(it["fmin"]) if it["fmin"] else "NA",
                    float(it["fmax"]) if it["fmax"] else "NA",
                    it["R"], it["L"], it["R_pol_total"], it["R_pol_win"], it["R_pol_out"],
                    it["maxRe"], it["maxIm"], it["rms_pct"], it["min_gamma"],
                    it["edge_fast"], it["edge_slow"], h4, lo, hi])
    for c, wd in zip("ABCDEFGHIJKLMNOPQRS",
                     (46, 11, 10, 11, 11, 11, 12, 15, 17, 18, 15, 15, 12, 12, 15, 15, 16, 17, 17)):
        ws2.column_dimensions[c].width = wd
    ws2.freeze_panes = "A2"

    ws3 = wb.create_sheet("Info")
    for row in [["protocol", PROTOCOL],
                ["axis mode", a.mode],
                ["assembly", note],
                ["tau axis", "%d points, %.6g .. %.6g s" % (len(axis), axis.min(), axis.max())],
                ["units", "tau in seconds, gamma in Ohm"],
                ["datasets", str(M)],
                ["filled / blank", "%d filled, %d blank" % (filled, data.size - filled)],
                ["note", a.note],
                ["", ""],
                ["BLANK CELLS", "A blank cell means that dataset has NO valid value at that tau"],
                ["", "(its measured window does not cover it). Blanks are NOT zeros."],
                ["WARNING", "tau outside a dataset's own measured window is EXTRAPOLATION."]]:
        ws3.append(row)
    ws3.column_dimensions["A"].width = 16
    ws3.column_dimensions["B"].width = 96
    ws3["A1"].font = Font(bold=True)

    wb.save(xlsx_path)
    print("saved", xlsx_path)


if __name__ == "__main__":
    main()
